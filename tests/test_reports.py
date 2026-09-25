from io import BytesIO
from datetime import datetime, timezone

import pytest
from openpyxl import load_workbook
from sqlalchemy import Boolean, DateTime, Integer, String

from app.dependencies import get_current_user
from app.main import app
from app.models.organization import Organization
from app.models.user import User

from tests.conftest import (
    TestingSessionLocal,
    client,
)


# -------------------------------------------------
# TEST DATA HELPERS
# -------------------------------------------------

def create_model(
    db,
    model,
    tag: str,
    **provided_values
):
    values = {}

    columns = {
        column.name: column
        for column in model.__table__.columns
    }

    for key, value in provided_values.items():
        if key in columns:
            values[key] = value

    for column in model.__table__.columns:
        if column.name in values:
            continue

        if (
            column.primary_key
            and column.autoincrement
        ):
            continue

        if column.nullable:
            continue

        if column.default is not None:
            continue

        if column.server_default is not None:
            continue

        if isinstance(column.type, String):
            if "email" in column.name:
                values[column.name] = (
                    f"{tag}-{column.name}@example.com"
                )

            elif "password" in column.name:
                values[column.name] = (
                    "test-password-hash"
                )

            else:
                values[column.name] = (
                    f"{column.name}-{tag}"
                )

        elif isinstance(
            column.type,
            Boolean
        ):
            values[column.name] = True

        elif isinstance(
            column.type,
            DateTime
        ):
            values[column.name] = (
                datetime.now(
                    timezone.utc
                )
            )

        elif isinstance(
            column.type,
            Integer
        ):
            values[column.name] = 1

    obj = model(**values)

    db.add(obj)
    db.flush()

    return obj


def create_organization(
    db,
    tag: str
):
    return create_model(
        db,
        Organization,
        tag,
        name=f"Organization {tag}",
        slug=f"organization-{tag}",
        is_active=True,
    )


def create_user(
    db,
    organization,
    tag: str,
    role: str
):
    return create_model(
        db,
        User,
        tag,
        organization_id=organization.id,
        name=f"{role.title()} {tag}",
        full_name=f"{role.title()} {tag}",
        email=f"{role}-{tag}@example.com",
        hashed_password="test-password-hash",
        password_hash="test-password-hash",
        role=role,
        is_active=True,
    )


def authenticate_as(user):
    app.dependency_overrides[
        get_current_user
    ] = lambda: user


@pytest.fixture(autouse=True)
def cleanup_user_override():
    yield

    app.dependency_overrides.pop(
        get_current_user,
        None,
    )


# -------------------------------------------------
# WEEKLY REPORT
# -------------------------------------------------

def test_admin_can_download_weekly_report():
    db = TestingSessionLocal()

    try:
        organization = create_organization(
            db,
            "weekly-report",
        )

        admin = create_user(
            db,
            organization,
            "weekly-admin",
            "admin",
        )

        db.commit()

        authenticate_as(admin)

        response = client.get(
            "/reports/operations.xlsx",
            params={
                "period": "weekly"
            },
        )

        assert response.status_code == 200

        assert (
            response.headers[
                "content-type"
            ]
            ==
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

        assert (
            "resolveops_weekly_report_"
            in response.headers[
                "content-disposition"
            ]
        )

        workbook = load_workbook(
            BytesIO(response.content)
        )

        assert workbook.sheetnames == [
            "Executive Summary",
            "Incidents",
            "SLA Performance",
            "Team Workload",
        ]

        summary = workbook[
            "Executive Summary"
        ]

        assert (
            summary["A1"].value
            ==
            "ResolveOps Operations Report"
        )

        assert (
            summary["B3"].value
            == "Weekly"
        )

    finally:
        db.close()


# -------------------------------------------------
# MONTHLY REPORT
# -------------------------------------------------

def test_manager_can_download_monthly_report():
    db = TestingSessionLocal()

    try:
        organization = create_organization(
            db,
            "monthly-report",
        )

        manager = create_user(
            db,
            organization,
            "monthly-manager",
            "manager",
        )

        db.commit()

        authenticate_as(manager)

        response = client.get(
            "/reports/operations.xlsx",
            params={
                "period": "monthly"
            },
        )

        assert response.status_code == 200

        assert (
            "resolveops_monthly_report_"
            in response.headers[
                "content-disposition"
            ]
        )

        workbook = load_workbook(
            BytesIO(response.content)
        )

        summary = workbook[
            "Executive Summary"
        ]

        assert (
            summary["B3"].value
            == "Monthly"
        )

    finally:
        db.close()


# -------------------------------------------------
# ALL-TIME REPORT
# -------------------------------------------------

def test_admin_can_download_all_time_report():
    db = TestingSessionLocal()

    try:
        organization = create_organization(
            db,
            "all-time-report",
        )

        admin = create_user(
            db,
            organization,
            "all-time-admin",
            "admin",
        )

        db.commit()

        authenticate_as(admin)

        response = client.get(
            "/reports/operations.xlsx",
            params={
                "period": "all_time"
            },
        )

        assert response.status_code == 200

        assert (
            "resolveops_all_time_report_"
            in response.headers[
                "content-disposition"
            ]
        )

        workbook = load_workbook(
            BytesIO(response.content)
        )

        assert workbook.sheetnames == [
            "Executive Summary",
            "Incidents",
            "SLA Performance",
            "Team Workload",
        ]

        summary = workbook[
            "Executive Summary"
        ]

        assert (
            summary["A1"].value
            ==
            "ResolveOps Operations Report"
        )

        assert (
            summary["B3"].value
            == "All Time"
        )

    finally:
        db.close()


# -------------------------------------------------
# AGENT CANNOT EXPORT MANAGEMENT REPORT
# -------------------------------------------------

def test_agent_cannot_download_report():
    db = TestingSessionLocal()

    try:
        organization = create_organization(
            db,
            "agent-report",
        )

        agent = create_user(
            db,
            organization,
            "report-agent",
            "agent",
        )

        db.commit()

        authenticate_as(agent)

        response = client.get(
            "/reports/operations.xlsx",
            params={
                "period": "weekly"
            },
        )

        assert response.status_code == 403

    finally:
        db.close()


# -------------------------------------------------
# CUSTOMER CANNOT EXPORT MANAGEMENT REPORT
# -------------------------------------------------

def test_customer_cannot_download_report():
    db = TestingSessionLocal()

    try:
        organization = create_organization(
            db,
            "customer-report",
        )

        customer = create_user(
            db,
            organization,
            "report-customer",
            "customer",
        )

        db.commit()

        authenticate_as(customer)

        response = client.get(
            "/reports/operations.xlsx",
            params={
                "period": "monthly"
            },
        )

        assert response.status_code == 403

    finally:
        db.close()


# -------------------------------------------------
# INVALID PERIOD
# -------------------------------------------------

def test_invalid_report_period_is_rejected():
    db = TestingSessionLocal()

    try:
        organization = create_organization(
            db,
            "invalid-period",
        )

        admin = create_user(
            db,
            organization,
            "invalid-admin",
            "admin",
        )

        db.commit()

        authenticate_as(admin)

        response = client.get(
            "/reports/operations.xlsx",
            params={
                "period": "yearly"
            },
        )

        assert response.status_code == 400

        assert response.json()["detail"] == (
            "Report period must be "
            "'weekly', 'monthly', or 'all_time'."
        )

    finally:
        db.close()
