from collections import Counter
from datetime import datetime, timedelta, timezone
from io import BytesIO

from fastapi import (
    APIRouter,
    Depends,
    HTTPException,
    Query,
    status,
)
from fastapi.responses import StreamingResponse

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter

from sqlalchemy.orm import Session

from app.database import get_db
from app.dependencies import require_roles

from app.models.incident import Incident
from app.models.user import User


router = APIRouter(
    prefix="/reports",
    tags=["Reports"],
)


# -------------------------------------------------
# DATE / VALUE HELPERS
# -------------------------------------------------

def as_utc(value):
    if value is None:
        return None

    if value.tzinfo is None:
        return value.replace(
            tzinfo=timezone.utc
        )

    return value.astimezone(
        timezone.utc
    )


def format_datetime(value):
    value = as_utc(value)

    if value is None:
        return ""

    return value.strftime(
        "%Y-%m-%d %H:%M:%S UTC"
    )


def safe_excel_value(value):
    """
    Prevent text values beginning with Excel formula
    characters from being interpreted as formulas.
    """

    if value is None:
        return ""

    if not isinstance(value, str):
        return value

    if value.startswith(
        ("=", "+", "-", "@")
    ):
        return f"'{value}"

    return value


def get_period_bounds(
    period: str,
    now: datetime
):
    if period == "weekly":
        start = (
            now
            - timedelta(days=6)
        ).replace(
            hour=0,
            minute=0,
            second=0,
            microsecond=0,
        )

        label = "Weekly"

    elif period == "monthly":
        start = now.replace(
            day=1,
            hour=0,
            minute=0,
            second=0,
            microsecond=0,
        )

        label = "Monthly"

    else:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=(
                "Report period must be "
                "'weekly', 'monthly', or 'all_time'."
            ),
        )

    return start, now, label


# -------------------------------------------------
# SLA HELPERS
# -------------------------------------------------

def get_response_sla_status(
    incident: Incident,
    now: datetime
):
    response_due = as_utc(
        incident.response_due_at
    )

    first_response = as_utc(
        incident.first_response_at
    )

    if response_due is None:
        return "Not configured"

    if first_response is None:
        if now > response_due:
            return "Breached"

        return "Pending"

    if first_response <= response_due:
        return "Met"

    return "Breached"


def get_resolution_sla_status(
    incident: Incident,
    now: datetime
):
    resolution_due = as_utc(
        incident.resolution_due_at
    )

    resolved_at = as_utc(
        incident.resolved_at
    )

    if resolution_due is None:
        return "Not configured"

    if resolved_at is None:
        if now > resolution_due:
            return "Breached"

        return "Pending"

    if resolved_at <= resolution_due:
        return "Met"

    return "Breached"


def get_overall_sla_status(
    response_status: str,
    resolution_status: str
):
    statuses = {
        response_status,
        resolution_status,
    }

    if "Breached" in statuses:
        return "Breached"

    if (
        response_status == "Met"
        and resolution_status == "Met"
    ):
        return "Met"

    if statuses == {"Not configured"}:
        return "Not configured"

    return "Pending"


def get_resolution_minutes(
    incident: Incident
):
    created_at = as_utc(
        incident.created_at
    )

    resolved_at = as_utc(
        incident.resolved_at
    )

    if (
        created_at is None
        or resolved_at is None
    ):
        return None

    minutes = (
        resolved_at - created_at
    ).total_seconds() / 60

    if minutes < 0:
        return None

    return round(minutes, 2)


# -------------------------------------------------
# EXCEL STYLING
# -------------------------------------------------

HEADER_FILL = PatternFill(
    fill_type="solid",
    fgColor="1E3A8A",
)

SECTION_FILL = PatternFill(
    fill_type="solid",
    fgColor="DBEAFE",
)

HEADER_FONT = Font(
    bold=True,
    color="FFFFFF",
)

TITLE_FONT = Font(
    bold=True,
    size=18,
)

SECTION_FONT = Font(
    bold=True,
    color="1E3A8A",
)

THIN_BORDER = Border(
    bottom=Side(
        style="thin",
        color="D1D5DB",
    )
)


def style_header_row(
    worksheet,
    row_number: int
):
    for cell in worksheet[row_number]:
        if cell.value is None:
            continue

        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(
            vertical="center"
        )


def auto_size_columns(
    worksheet,
    maximum_width=45
):
    for column_cells in worksheet.columns:
        maximum_length = 0

        column_letter = get_column_letter(
            column_cells[0].column
        )

        for cell in column_cells:
            try:
                value = (
                    ""
                    if cell.value is None
                    else str(cell.value)
                )

                maximum_length = max(
                    maximum_length,
                    len(value),
                )
            except Exception:
                continue

        worksheet.column_dimensions[
            column_letter
        ].width = min(
            maximum_length + 3,
            maximum_width,
        )


def style_data_rows(
    worksheet,
    start_row: int
):
    for row in worksheet.iter_rows(
        min_row=start_row
    ):
        for cell in row:
            cell.border = THIN_BORDER
            cell.alignment = Alignment(
                vertical="top"
            )


# -------------------------------------------------
# BUILD EXCEL WORKBOOK
# -------------------------------------------------

def build_operations_workbook(
    incidents,
    users,
    start_date,
    end_date,
    period_label,
):
    now = datetime.now(
        timezone.utc
    )

    workbook = Workbook()

    summary_sheet = workbook.active
    summary_sheet.title = (
        "Executive Summary"
    )

    incidents_sheet = (
        workbook.create_sheet(
            "Incidents"
        )
    )

    sla_sheet = (
        workbook.create_sheet(
            "SLA Performance"
        )
    )

    workload_sheet = (
        workbook.create_sheet(
            "Team Workload"
        )
    )


    # -------------------------------------------------
    # USER LOOKUPS
    # -------------------------------------------------

    user_lookup = {
        user.id: user
        for user in users
    }

    active_staff = [
        user
        for user in users
        if (
            user.is_active
            and user.role in {
                "admin",
                "manager",
                "agent",
            }
        )
    ]


    # -------------------------------------------------
    # KPI CALCULATIONS
    # -------------------------------------------------

    total_incidents = len(incidents)

    active_statuses = {
        "open",
        "assigned",
        "in_progress",
    }

    active_incidents = sum(
        1
        for incident in incidents
        if incident.status
        in active_statuses
    )

    status_counts = Counter(
        incident.status
        for incident in incidents
    )

    priority_counts = Counter(
        incident.priority
        for incident in incidents
    )

    category_counts = Counter(
        (
            incident.category.strip()
            if (
                incident.category
                and incident.category.strip()
            )
            else "uncategorized"
        )
        for incident in incidents
    )

    breached_incidents = 0

    completed_sla_incidents = 0
    compliant_sla_incidents = 0

    resolution_times = []

    sla_rows = []

    for incident in incidents:
        response_status = (
            get_response_sla_status(
                incident,
                now,
            )
        )

        resolution_status = (
            get_resolution_sla_status(
                incident,
                now,
            )
        )

        overall_status = (
            get_overall_sla_status(
                response_status,
                resolution_status,
            )
        )

        if overall_status == "Breached":
            breached_incidents += 1

        response_due = as_utc(
            incident.response_due_at
        )

        resolution_due = as_utc(
            incident.resolution_due_at
        )

        first_response = as_utc(
            incident.first_response_at
        )

        resolved_at = as_utc(
            incident.resolved_at
        )

        if (
            response_due is not None
            and resolution_due is not None
            and resolved_at is not None
        ):
            completed_sla_incidents += 1

            response_met = (
                first_response is not None
                and first_response
                <= response_due
            )

            resolution_met = (
                resolved_at
                <= resolution_due
            )

            if (
                response_met
                and resolution_met
            ):
                compliant_sla_incidents += 1

        resolution_minutes = (
            get_resolution_minutes(
                incident
            )
        )

        if resolution_minutes is not None:
            resolution_times.append(
                resolution_minutes
            )

        sla_rows.append(
            {
                "incident": incident,
                "response_status":
                    response_status,
                "resolution_status":
                    resolution_status,
                "overall_status":
                    overall_status,
                "resolution_minutes":
                    resolution_minutes,
            }
        )

    if completed_sla_incidents > 0:
        sla_compliance = round(
            (
                compliant_sla_incidents
                / completed_sla_incidents
            )
            * 100,
            2,
        )
    else:
        sla_compliance = 0.0

    if resolution_times:
        average_resolution = round(
            sum(resolution_times)
            / len(resolution_times),
            2,
        )
    else:
        average_resolution = None


    # -------------------------------------------------
    # EXECUTIVE SUMMARY SHEET
    # -------------------------------------------------

    summary_sheet.merge_cells(
        "A1:D1"
    )

    summary_sheet["A1"] = (
        "ResolveOps Operations Report"
    )

    summary_sheet["A1"].font = TITLE_FONT

    summary_sheet["A3"] = "Report Period"
    summary_sheet["B3"] = period_label

    summary_sheet["A4"] = "From"
    summary_sheet["B4"] = (
        format_datetime(start_date)
    )

    summary_sheet["A5"] = "To"
    summary_sheet["B5"] = (
        format_datetime(end_date)
    )

    summary_sheet["A7"] = (
        "Key Performance Indicators"
    )
    summary_sheet["A7"].font = SECTION_FONT
    summary_sheet["A7"].fill = SECTION_FILL

    summary_sheet.append(
        [
            "Metric",
            "Value",
        ]
    )

    style_header_row(
        summary_sheet,
        8,
    )

    summary_sheet.append(
        [
            "Total Incidents",
            total_incidents,
        ]
    )

    summary_sheet.append(
        [
            "Active Incidents",
            active_incidents,
        ]
    )

    summary_sheet.append(
        [
            "Breached Incidents",
            breached_incidents,
        ]
    )

    summary_sheet.append(
        [
            "SLA Compliance",
            f"{sla_compliance}%",
        ]
    )

    summary_sheet.append(
        [
            "Average Resolution",
            (
                f"{average_resolution} min"
                if average_resolution
                is not None
                else "N/A"
            ),
        ]
    )

    summary_sheet.append(
        [
            "Completed SLA Incidents",
            completed_sla_incidents,
        ]
    )


    # STATUS BREAKDOWN

    summary_sheet["D7"] = (
        "Status Breakdown"
    )
    summary_sheet["D7"].font = SECTION_FONT
    summary_sheet["D7"].fill = SECTION_FILL

    summary_sheet["D8"] = "Status"
    summary_sheet["E8"] = "Count"

    for cell in (
        summary_sheet["D8"],
        summary_sheet["E8"],
    ):
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    status_order = [
        "open",
        "assigned",
        "in_progress",
        "resolved",
        "closed",
    ]

    row_number = 9

    for status_name in status_order:
        summary_sheet.cell(
            row=row_number,
            column=4,
            value=status_name.replace(
                "_",
                " ",
            ).title(),
        )

        summary_sheet.cell(
            row=row_number,
            column=5,
            value=status_counts.get(
                status_name,
                0,
            ),
        )

        row_number += 1


    # PRIORITY BREAKDOWN

    summary_sheet["G7"] = (
        "Priority Breakdown"
    )
    summary_sheet["G7"].font = SECTION_FONT
    summary_sheet["G7"].fill = SECTION_FILL

    summary_sheet["G8"] = "Priority"
    summary_sheet["H8"] = "Count"

    for cell in (
        summary_sheet["G8"],
        summary_sheet["H8"],
    ):
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    priority_order = [
        "critical",
        "high",
        "medium",
        "low",
    ]

    row_number = 9

    for priority_name in priority_order:
        summary_sheet.cell(
            row=row_number,
            column=7,
            value=priority_name.title(),
        )

        summary_sheet.cell(
            row=row_number,
            column=8,
            value=priority_counts.get(
                priority_name,
                0,
            ),
        )

        row_number += 1


    # CATEGORY BREAKDOWN

    summary_sheet["J7"] = (
        "Category Breakdown"
    )
    summary_sheet["J7"].font = SECTION_FONT
    summary_sheet["J7"].fill = SECTION_FILL

    summary_sheet["J8"] = "Category"
    summary_sheet["K8"] = "Count"

    for cell in (
        summary_sheet["J8"],
        summary_sheet["K8"],
    ):
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    row_number = 9

    for category, count in sorted(
        category_counts.items(),
        key=lambda item: (
            -item[1],
            item[0].lower(),
        ),
    ):
        summary_sheet.cell(
            row=row_number,
            column=10,
            value=safe_excel_value(
                category.title()
            ),
        )

        summary_sheet.cell(
            row=row_number,
            column=11,
            value=count,
        )

        row_number += 1

    summary_sheet.freeze_panes = "A8"

    auto_size_columns(
        summary_sheet
    )


    # -------------------------------------------------
    # INCIDENTS SHEET
    # -------------------------------------------------

    incident_headers = [
        "Ticket",
        "Title",
        "Priority",
        "Category",
        "Status",
        "Created By",
        "Assigned To",
        "Created At (UTC)",
        "First Response At (UTC)",
        "Resolved At (UTC)",
        "Response Due (UTC)",
        "Resolution Due (UTC)",
        "Response SLA",
        "Resolution SLA",
        "Overall SLA",
        "Resolution Minutes",
    ]

    incidents_sheet.append(
        incident_headers
    )

    style_header_row(
        incidents_sheet,
        1,
    )

    for sla_row in sorted(
        sla_rows,
        key=lambda item: (
            as_utc(
                item["incident"].created_at
            )
            or datetime.min.replace(
                tzinfo=timezone.utc
            )
        ),
        reverse=True,
    ):
        incident = sla_row["incident"]

        creator = user_lookup.get(
            incident.created_by
        )

        assignee = user_lookup.get(
            incident.assigned_to
        )

        incidents_sheet.append(
            [
                safe_excel_value(
                    incident.ticket_number
                ),
                safe_excel_value(
                    incident.title
                ),
                safe_excel_value(
                    incident.priority.title()
                    if incident.priority
                    else ""
                ),
                safe_excel_value(
                    incident.category
                    or "Uncategorized"
                ),
                safe_excel_value(
                    incident.status.replace(
                        "_",
                        " ",
                    ).title()
                    if incident.status
                    else ""
                ),
                safe_excel_value(
                    creator.name
                    if creator
                    else ""
                ),
                safe_excel_value(
                    assignee.name
                    if assignee
                    else "Unassigned"
                ),
                format_datetime(
                    incident.created_at
                ),
                format_datetime(
                    incident.first_response_at
                ),
                format_datetime(
                    incident.resolved_at
                ),
                format_datetime(
                    incident.response_due_at
                ),
                format_datetime(
                    incident.resolution_due_at
                ),
                sla_row[
                    "response_status"
                ],
                sla_row[
                    "resolution_status"
                ],
                sla_row[
                    "overall_status"
                ],
                sla_row[
                    "resolution_minutes"
                ],
            ]
        )

    incidents_sheet.freeze_panes = "A2"
    incidents_sheet.auto_filter.ref = (
        incidents_sheet.dimensions
    )

    style_data_rows(
        incidents_sheet,
        2,
    )

    auto_size_columns(
        incidents_sheet
    )


    # -------------------------------------------------
    # SLA PERFORMANCE SHEET
    # -------------------------------------------------

    sla_headers = [
        "Ticket",
        "Priority",
        "Response Due (UTC)",
        "First Response (UTC)",
        "Response SLA",
        "Resolution Due (UTC)",
        "Resolved At (UTC)",
        "Resolution SLA",
        "Overall SLA",
    ]

    sla_sheet.append(
        sla_headers
    )

    style_header_row(
        sla_sheet,
        1,
    )

    for sla_row in sla_rows:
        incident = sla_row["incident"]

        sla_sheet.append(
            [
                safe_excel_value(
                    incident.ticket_number
                ),
                safe_excel_value(
                    incident.priority.title()
                    if incident.priority
                    else ""
                ),
                format_datetime(
                    incident.response_due_at
                ),
                format_datetime(
                    incident.first_response_at
                ),
                sla_row[
                    "response_status"
                ],
                format_datetime(
                    incident.resolution_due_at
                ),
                format_datetime(
                    incident.resolved_at
                ),
                sla_row[
                    "resolution_status"
                ],
                sla_row[
                    "overall_status"
                ],
            ]
        )

    sla_sheet.freeze_panes = "A2"
    sla_sheet.auto_filter.ref = (
        sla_sheet.dimensions
    )

    style_data_rows(
        sla_sheet,
        2,
    )

    auto_size_columns(
        sla_sheet
    )


    # -------------------------------------------------
    # TEAM WORKLOAD SHEET
    # -------------------------------------------------

    assigned_counts = Counter(
        incident.assigned_to
        for incident in incidents
        if incident.assigned_to
        is not None
    )

    active_assigned_counts = Counter(
        incident.assigned_to
        for incident in incidents
        if (
            incident.assigned_to
            is not None
            and incident.status
            in {
                "assigned",
                "in_progress",
            }
        )
    )

    workload_sheet.append(
        [
            "Team Member",
            "Role",
            "Assigned Incidents",
            "Active Incidents",
        ]
    )

    style_header_row(
        workload_sheet,
        1,
    )

    for user in sorted(
        active_staff,
        key=lambda item: (
            -active_assigned_counts.get(
                item.id,
                0,
            ),
            item.name.lower(),
        ),
    ):
        workload_sheet.append(
            [
                safe_excel_value(
                    user.name
                ),
                safe_excel_value(
                    user.role.title()
                ),
                assigned_counts.get(
                    user.id,
                    0,
                ),
                active_assigned_counts.get(
                    user.id,
                    0,
                ),
            ]
        )

    workload_sheet.freeze_panes = "A2"
    workload_sheet.auto_filter.ref = (
        workload_sheet.dimensions
    )

    style_data_rows(
        workload_sheet,
        2,
    )

    auto_size_columns(
        workload_sheet
    )


    # -------------------------------------------------
    # SAVE WORKBOOK TO MEMORY
    # -------------------------------------------------

    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    return output


# -------------------------------------------------
# DOWNLOAD OPERATIONS REPORT
# -------------------------------------------------

@router.get(
    "/operations.xlsx"
)
def download_operations_report(
    period: str = Query(
        default="weekly"
    ),
    current_user: User = Depends(
        require_roles(
            "admin",
            "manager",
        )
    ),
    db: Session = Depends(get_db),
):
    period = period.lower().strip()

    now = datetime.now(
        timezone.utc
    )

    all_incidents = (
        db.query(Incident)
        .filter(
            Incident.organization_id
            == current_user.organization_id
        )
        .all()
    )

    # -------------------------------------------------
    # REPORT PERIOD
    #
    # weekly   -> last 7 calendar days including today
    # monthly  -> current calendar month
    # all_time -> every incident stored for the tenant
    # -------------------------------------------------

    if period == "all_time":
        incidents = list(all_incidents)

        created_dates = [
            as_utc(incident.created_at)
            for incident in all_incidents
            if incident.created_at is not None
        ]

        start_date = (
            min(created_dates)
            if created_dates
            else now
        )

        end_date = now
        period_label = "All Time"

    else:
        start_date, end_date, period_label = (
            get_period_bounds(
                period,
                now,
            )
        )

        incidents = []

        for incident in all_incidents:
            created_at = as_utc(
                incident.created_at
            )

            if created_at is None:
                continue

            if (
                start_date
                <= created_at
                <= end_date
            ):
                incidents.append(
                    incident
                )

    users = (
        db.query(User)
        .filter(
            User.organization_id
            == current_user.organization_id
        )
        .all()
    )

    workbook = build_operations_workbook(
        incidents=incidents,
        users=users,
        start_date=start_date,
        end_date=end_date,
        period_label=period_label,
    )

    filename = (
        f"resolveops_"
        f"{period}_report_"
        f"{now.date().isoformat()}.xlsx"
    )

    return StreamingResponse(
        workbook,
        media_type=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition":
                f'attachment; filename="{filename}"'
        },
    )